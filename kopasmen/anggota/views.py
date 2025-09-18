from urllib import request
from django.shortcuts import render, redirect, get_object_or_404
from django.core.paginator import Paginator
from admin_koperasi.models import Admin
from .models import Anggota
from .forms import AdminForm, AnggotaForm
from django.db import connection

def kelola_akun_view(request):
    anggotas = Anggota.objects.all() 
    admins = Admin.objects.all()
    return render(request, 'kelola_akun.html', {'anggotas': anggotas, 'admins': admins})

def kelola_akun(request):
    admins = Admin.objects.all().order_by("id_admin")
    anggotas = Anggota.objects.all().order_by("nomor_anggota")

    # Pagination
    admin_paginator = Paginator(admins, 20)   # 20 per halaman
    anggota_paginator = Paginator(anggotas, 20)

    admin_page_number = request.GET.get("page_admin")
    anggota_page_number = request.GET.get("page_anggota")

    admin_page = admin_paginator.get_page(admin_page_number)
    anggota_page = anggota_paginator.get_page(anggota_page_number)

    return render(request, "kelola_akun.html", {
        "admins": admin_page,
        "anggotas": anggota_page,
    })

def tambah_admin(request):
    if request.method == 'POST':
        form = AdminForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('kelola_akun') 
    else:
        form = AdminForm()
    return render(request, 'form_admin.html', {'form': form})

def edit_admin(request, id_admin):
    admin = get_object_or_404(Admin, id_admin=id_admin)
    form = AdminForm(request.POST or None, instance=admin)
    if form.is_valid():
        form.save()
        return redirect('kelola_akun')
    return render(request, 'form_admin.html', {'form': form, 'judul': 'Edit Admin'})

def hapus_admin(request, id_admin):
    admin = get_object_or_404(Admin, id_admin=id_admin)
    admin.delete()
    return redirect('kelola_akun')

def detail_admin(request, id_admin):
    admin = get_object_or_404(Admin, id_admin=id_admin)
    all_admins = Admin.objects.order_by('id_admin')
    nomor_urut = list(all_admins).index(admin) + 1  

    context = {
        'admin': admin,
        'nomor_urut': nomor_urut,
    }
    return render(request, 'detailA.html', context)

def tambah_anggota(request):
    if request.method == 'POST':
        form = AnggotaForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('kelola_akun')
    else:
        form = AnggotaForm() 
    return render(request, 'form_admin.html', {'form': form, 'judul': 'Tambah Anggota'})


def anggota_detail(request, nomor_anggota):
    try:
        anggota = Anggota.objects.get(nomor_anggota=nomor_anggota)
        return render(request, 'detail.html', {'anggota': anggota})
    except Anggota.DoesNotExist:
        return render(request, 'detail.html', {'error': 'Anggota not found'})
    
def edit_anggota(request, nomor_anggota):
    anggota = get_object_or_404(Anggota, nomor_anggota=nomor_anggota)
    form = AnggotaForm(request.POST or None, instance=anggota)
    if form.is_valid():
        form.save()
        return redirect('kelola_akun')
    return render(request, 'form_anggota.html', {'form': form, 'judul': 'Edit Anggota'})


def hapus_anggota(request, nomor_anggota):
    anggota = get_object_or_404(Anggota, nomor_anggota=nomor_anggota)
    anggota.delete()
    return redirect('kelola_akun')

def detail_anggota(request, nomor_anggota):
    anggota = get_object_or_404(Anggota, nomor_anggota=nomor_anggota)
    return render(request, 'detail.html', {'anggota': anggota})

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.http import HttpResponse
from .models import Anggota

def export_excel_anggota(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daftar Anggota"


    headers = ["No. Anggota", "Nama", "NIP", "Alamat", "No. Telepon", "Email", "Jenis Kelamin", "Tanggal Daftar", "Status"]
    ws.append(headers)


    header_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")  # kuning
    header_font = Font(bold=True, color="000000")  # teks hitam tebal
    header_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(border_style="thin", color="000000"),
        right=Side(border_style="thin", color="000000"),
        top=Side(border_style="thin", color="000000"),
        bottom=Side(border_style="thin", color="000000")
    )

    for col in ws.iter_cols(min_row=1, max_row=1, min_col=1, max_col=len(headers)):
        for cell in col:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border


    for anggota in Anggota.objects.all():
        ws.append([
            anggota.nomor_anggota,
            anggota.nama,
            anggota.nip,
            anggota.alamat,
            anggota.no_telp,
            anggota.email,
            anggota.jenis_kelamin,
            anggota.tanggal_daftar.strftime("%d-%m-%Y"),
            anggota.status
        ])

   
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top")

    
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_length + 2

    # Response
    response = HttpResponse(content_type="application/ms-excel")
    response["Content-Disposition"] = 'attachment; filename="anggota koperasi.xlsx"'
    wb.save(response)
    return response



from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from django.http import HttpResponse
from .models import Anggota

def export_pdf_anggota(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="anggota koperasi.pdf"'

    
    doc = SimpleDocTemplate(response, pagesize=A4)
    elements = []

    styles = getSampleStyleSheet()
    normal_style = styles["Normal"]  
    normal_style.fontSize = 7

    title = Paragraph("Daftar Anggota", styles['Title'])
    elements.append(title)

    
    data = [["No. Anggota", "Nama", "NIP", "Alamat", "No. Telepon", 
             "Email", "Jenis Kelamin", "Tanggal Daftar", "Status"]]

    
    for anggota in Anggota.objects.all():
        data.append([
            Paragraph(str(anggota.nomor_anggota), normal_style),
            Paragraph(anggota.nama, normal_style),
            Paragraph(anggota.nip, normal_style),
            Paragraph(anggota.alamat or "", normal_style),
            Paragraph(anggota.no_telp or "", normal_style),
            Paragraph(anggota.email or "", normal_style),
            Paragraph(anggota.jenis_kelamin, normal_style),
            Paragraph(anggota.tanggal_daftar.strftime("%d-%m-%Y"), normal_style),
            Paragraph(anggota.status, normal_style)
        ])

    
    col_widths = [50, 75, 60, 85, 70, 80, 50, 60, 50]

    table = Table(data, repeatRows=1, colWidths=col_widths)
    table.setStyle(TableStyle([
        ("FONTSIZE", (0,0), (-1,-1), 7),
        ("BACKGROUND", (0,0), (-1,0), colors.yellow),
        ("TEXTCOLOR", (0,0), (-1,0), colors.black),
        ("ALIGN", (0,0), (-1,0), "CENTER"),  
        ("ALIGN", (0,1), (-1,-1), "LEFT"),   
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("BOTTOMPADDING", (0,0), (-1,0), 5),
        ("GRID", (0,0), (-1,-1), 0.5, colors.grey),
    ]))

    elements.append(table)
    doc.build(elements)
    return response

from django.shortcuts import render, redirect
from django.db.models import Sum
from admin_koperasi.models import Admin  # import dari app admin_koperasi
from .models import Anggota  # model anggota ada di app ini
from django.contrib import messages

def kelola_akun_view(request):
    # cek session login
    if not request.session.get('admin_id'):
        return redirect('admin_koperasi:login')

    role = request.session.get('admin_role')
    username = request.session.get('admin_username')

    admins = Admin.objects.all()
    anggotas = Anggota.objects.all()

    context = {
        'username': username,
        'role': role,
        'admins': admins,
        'anggotas': anggotas,
    }

    return render(request, 'kelola_akun.html', context)

def tambah_admin(request):
    if not request.session.get('admin_id'):
        return redirect('admin_koperasi:login')

    role = request.session.get('admin_role')
    username = request.session.get('admin_username')

    if request.method == 'POST':
        form = AdminForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Admin berhasil ditambahkan.")
            return redirect('kelola_akun')
    else:
        form = AdminForm()

    context = {
        'form': form,
        'judul': 'Tambah Admin',
        'username': username,
        'role': role,
    }
    return render(request, 'form_admin.html', context)

def tambah_anggota(request):
    if not request.session.get('admin_id'):
        return redirect('admin_koperasi:login')

    role = request.session.get('admin_role')
    username = request.session.get('admin_username')

    if request.method == 'POST':
        form = AnggotaForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Anggota berhasil ditambahkan.")
            return redirect('kelola_akun')
    else:
        form = AnggotaForm()

    context = {
        'form': form,
        'judul': 'Tambah Anggota',
        'username': username,
        'role': role,
    }
    return render(request, 'form_admin.html', context)

def edit_admin(request, id_admin):
    if not request.session.get('admin_id'):
        return redirect('admin_koperasi:login')

    role = request.session.get('admin_role')
    username = request.session.get('admin_username')

    admin = get_object_or_404(Admin, id_admin=id_admin)
    form = AdminForm(request.POST or None, instance=admin)
    if form.is_valid():
        form.save()
        messages.success(request, "Admin berhasil diperbarui.")
        return redirect('kelola_akun')

    context = {
        'form': form,
        'judul': 'Edit Admin',
        'username': username,
        'role': role,
    }
    return render(request, 'form_admin.html', context)

def edit_anggota(request, nomor_anggota):
    if not request.session.get('admin_id'):
        return redirect('admin_koperasi:login')

    role = request.session.get('admin_role')
    username = request.session.get('admin_username')

    anggota = get_object_or_404(Anggota, nomor_anggota=nomor_anggota)
    form = AnggotaForm(request.POST or None, instance=anggota)
    if form.is_valid():
        form.save()
        messages.success(request, "Anggota berhasil diperbarui.")
        return redirect('kelola_akun')

    context = {
        'form': form,
        'judul': 'Edit Anggota',
        'username': username,
        'role': role,
    }
    return render(request, 'form_anggota.html', context)

def detail_admin(request, id_admin):
    if not request.session.get('admin_id'):
        return redirect('admin_koperasi:login')

    role = request.session.get('admin_role')
    username = request.session.get('admin_username')

    admin = get_object_or_404(Admin, id_admin=id_admin)
    all_admins = Admin.objects.order_by('id_admin')
    nomor_urut = list(all_admins).index(admin) + 1  

    context = {
        'admin': admin,
        'nomor_urut': nomor_urut,
        'username': username,
        'role': role,
    }
    return render(request, 'detailA.html', context)

def detail_anggota(request, nomor_anggota):
    if not request.session.get('admin_id'):
        return redirect('admin_koperasi:login')

    role = request.session.get('admin_role')
    username = request.session.get('admin_username')

    anggota = get_object_or_404(Anggota, nomor_anggota=nomor_anggota)

    context = {
        'anggota': anggota,
        'username': username,
        'role': role,
    }
    return render(request, 'detail.html', context)



